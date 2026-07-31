from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE = Path(__file__).resolve().parents[1]
XLSX_PATH = BASE / "Расчет_вероятности_захвата_нейтронов_B4C.xlsx"
PNG_PATH = BASE / "Расчет_вероятности_захвата_нейтронов_B4C.png"


def build_formula(thickness_cell: str, enrichment_cell: str) -> str:
    return (
        "=1-EXP(-((4*Параметры!$B$2/Параметры!$B$3*Параметры!$B$10*"
        f"{enrichment_cell})*(Параметры!$B$6*Параметры!$B$7/Параметры!$B$8*1E-24))"
        f"*({thickness_cell}*1E-7)/SIN(RADIANS(Параметры!$B$9)))"
    )


def create_workbook() -> None:
    wb = Workbook()
    ws_params = wb.active
    ws_params.title = "Параметры"
    ws_calc = wb.create_sheet("Расчет")
    ws_chart = wb.create_sheet("График")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    subheader_fill = PatternFill("solid", fgColor="EAF3E8")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font = Font(name="Arial", size=12)
    header_font = Font(name="Arial", size=12, bold=True)
    title_font = Font(name="Arial", size=14, bold=True)

    params = [
        ("Плотность B₄C", 2.52, "г/см³"),
        ("Молярная масса B₄C", 55.26, "г/моль"),
        ("Доля ¹⁰B для природного бора", 0.20, "доля"),
        ("Доля ¹⁰B для обогащённого бора", 0.90, "доля"),
        ("Сечение захвата при 1.8 Å", 3840, "барн"),
        ("Длина волны нейтронов", 1.8, "Å"),
        ("Нормировочная длина волны", 1.8, "Å"),
        ("Угол между пучком и поверхностью", 90, "градусы"),
        ("Постоянная Авогадро", 6.02214076e23, "моль⁻¹"),
    ]

    for col, value in enumerate(["Параметр", "Значение", "Единица"], 1):
        cell = ws_params.cell(1, col, value)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for row, (name, value, unit) in enumerate(params, 2):
        ws_params.cell(row, 1, name)
        ws_params.cell(row, 2, value)
        ws_params.cell(row, 3, unit)
        for col in range(1, 4):
            ws_params.cell(row, col).font = font
            ws_params.cell(row, col).border = border
        ws_params.cell(row, 2).number_format = "0.00E+00" if abs(value) > 1e6 else "0.########"

    ws_params["E1"] = "Расчётные величины"
    ws_params["E1"].font = title_font
    derived = [
        ("sigma_lambda, cm²", "=B6*B7/B8*1E-24"),
        ("n_10B, природный бор, см⁻³", "=4*B2/B3*B10*B4"),
        ("n_10B, обогащённый бор, см⁻³", "=4*B2/B3*B10*B5"),
        ("Sigma, природный бор, см⁻¹", "=F3*F2"),
        ("Sigma, обогащённый бор, см⁻¹", "=F4*F2"),
    ]
    for row, (name, formula) in enumerate(derived, 2):
        ws_params.cell(row, 5, name)
        ws_params.cell(row, 6, formula)
        for col in (5, 6):
            ws_params.cell(row, col).font = font
            ws_params.cell(row, col).border = border
        ws_params.cell(row, 6).number_format = "0.000E+00"

    for col, width in {"A": 34, "B": 18, "C": 16, "E": 34, "F": 20}.items():
        ws_params.column_dimensions[col].width = width
    ws_params.freeze_panes = "A2"

    headers = [
        "Толщина покрытия B₄C, нм",
        "Вероятность захвата, природный бор",
        "Вероятность захвата, обогащённый бор",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws_calc.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row, thickness_nm in enumerate(range(0, 3001, 10), 2):
        ws_calc.cell(row, 1, thickness_nm)
        ws_calc.cell(row, 2, build_formula(f"A{row}", "Параметры!$B$4"))
        ws_calc.cell(row, 3, build_formula(f"A{row}", "Параметры!$B$5"))
        for col in range(1, 4):
            ws_calc.cell(row, col).font = font
            ws_calc.cell(row, col).border = border
        ws_calc.cell(row, 1).number_format = "0"
        ws_calc.cell(row, 2).number_format = "0.0000"
        ws_calc.cell(row, 3).number_format = "0.0000"

    sample_headers = ["Образец", "Номинальная толщина, нм", "P, природный бор", "P, обогащённый бор"]
    for col, header in enumerate(sample_headers, 5):
        cell = ws_calc.cell(1, col, header)
        cell.font = header_font
        cell.fill = subheader_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    samples = [("B₄C-50", 50), ("B₄C-100", 100), ("B₄C-200", 200), ("B₄C-400", 400)]
    for row, (name, thickness_nm) in enumerate(samples, 2):
        ws_calc.cell(row, 5, name)
        ws_calc.cell(row, 6, thickness_nm)
        ws_calc.cell(row, 7, build_formula(f"F{row}", "Параметры!$B$4"))
        ws_calc.cell(row, 8, build_formula(f"F{row}", "Параметры!$B$5"))
        for col in range(5, 9):
            ws_calc.cell(row, col).font = font
            ws_calc.cell(row, col).border = border
        ws_calc.cell(row, 7).number_format = "0.0000"
        ws_calc.cell(row, 8).number_format = "0.0000"

    for col, width in {"A": 24, "B": 34, "C": 36, "E": 16, "F": 24, "G": 20, "H": 22}.items():
        ws_calc.column_dimensions[col].width = width
    ws_calc.freeze_panes = "A2"

    calc_table = Table(displayName="Tbl_Capture_Curve", ref="A1:C302")
    calc_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws_calc.add_table(calc_table)
    sample_table = Table(displayName="Tbl_Samples", ref="E1:H5")
    sample_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    ws_calc.add_table(sample_table)

    ws_chart["A1"] = "P(d) = 1 - exp[-Σd / sin(θ)]"
    ws_chart["A1"].font = Font(name="Arial", size=16, bold=True)
    ws_chart["A3"] = (
        "Расчёт выполнен по экспоненциальному закону ослабления нейтронного пучка. "
        "График показывает вероятность захвата нейтрона в слое B₄C, а не полную эффективность "
        "регистрации детектора. Прямые нейтронные измерения исследуемых образцов не проводились."
    )
    ws_chart["A3"].font = font
    ws_chart["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_chart.merge_cells("A3:H5")

    ws_chart["J1"] = "Данные вертикальных маркеров"
    ws_chart["J1"].font = header_font
    for idx, (name, thickness_nm) in enumerate(samples):
        col = 10 + idx * 2
        ws_chart.cell(2, col, f"{name} X")
        ws_chart.cell(2, col + 1, f"{name} Y")
        ws_chart.cell(3, col, thickness_nm)
        ws_chart.cell(3, col + 1, 0)
        ws_chart.cell(4, col, thickness_nm)
        ws_chart.cell(4, col + 1, 1)
        for row in range(2, 5):
            for c in (col, col + 1):
                ws_chart.cell(row, c).font = font
                ws_chart.cell(row, c).border = border

    ws_chart["J7"] = "Исследуемые образцы"
    ws_chart["J7"].font = header_font
    for col, header in enumerate(sample_headers, 10):
        cell = ws_chart.cell(8, col, header)
        cell.font = header_font
        cell.fill = subheader_fill
        cell.border = border
    for source_row in range(2, 6):
        target_row = source_row + 7
        for source_col in range(5, 9):
            target_col = source_col + 5
            cell = ws_chart.cell(target_row, target_col, f"=Расчет!{get_column_letter(source_col)}{source_row}")
            cell.font = font
            cell.border = border
            if source_col in (7, 8):
                cell.number_format = "0.0000"

    for col in range(1, 14):
        ws_chart.column_dimensions[get_column_letter(col)].width = 16 if col < 10 else 22
    ws_chart.sheet_view.showGridLines = False

    chart = ScatterChart()
    chart.title = "Расчётная зависимость вероятности захвата нейтрона от толщины покрытия B₄C"
    chart.style = 2
    chart.x_axis.title = "Толщина покрытия B₄C, нм"
    chart.y_axis.title = "Расчётная вероятность захвата нейтрона"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = 3000
    chart.legend.position = "r"
    chart.height = 16
    chart.width = 28

    x_values = Reference(ws_calc, min_col=1, min_row=2, max_row=302)
    for y_col, title, color in [
        (2, "Природный бор", "1F77B4"),
        (3, "Обогащённый бор", "D62728"),
    ]:
        y_values = Reference(ws_calc, min_col=y_col, min_row=2, max_row=302)
        series = Series(y_values, x_values, title=title)
        series.smooth = True
        series.marker.symbol = "none"
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 22000
        chart.series.append(series)

    for idx, (name, _thickness_nm) in enumerate(samples):
        col = 10 + idx * 2
        x_ref = Reference(ws_chart, min_col=col, min_row=3, max_row=4)
        y_ref = Reference(ws_chart, min_col=col + 1, min_row=3, max_row=4)
        series = Series(y_ref, x_ref, title=name)
        series.marker.symbol = "none"
        series.graphicalProperties.line.solidFill = "666666"
        series.graphicalProperties.line.width = 12000
        series.graphicalProperties.line.dashStyle = "dash"
        chart.series.append(series)

    ws_chart.add_chart(chart, "A7")

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(XLSX_PATH)


def create_png() -> None:
    rho = 2.52
    molar_mass = 55.26
    natural_enrichment = 0.20
    enriched_enrichment = 0.90
    sigma_0 = 3840
    wavelength = 1.8
    wavelength_0 = 1.8
    theta_deg = 90
    avogadro = 6.02214076e23
    barn = 1e-24
    nm_to_cm = 1e-7

    def probability(d_nm: np.ndarray, enrichment: float) -> np.ndarray:
        n_10b = 4 * rho / molar_mass * avogadro * enrichment
        sigma_lambda = sigma_0 * wavelength / wavelength_0 * barn
        sigma = n_10b * sigma_lambda
        return 1 - np.exp(-sigma * (d_nm * nm_to_cm) / np.sin(np.deg2rad(theta_deg)))

    d = np.arange(0, 3000 + 10, 10)
    y_natural = probability(d, natural_enrichment)
    y_enriched = probability(d, enriched_enrichment)
    samples = [("B$_4$C-50", 50), ("B$_4$C-100", 100), ("B$_4$C-200", 200), ("B$_4$C-400", 400)]

    plt.rcParams.update({"font.family": "Arial", "font.size": 14})
    fig, ax = plt.subplots(figsize=(19.2, 10.8), dpi=100)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    ax.plot(d, y_natural, label="Природный бор", linewidth=3, color="#1f77b4")
    ax.plot(d, y_enriched, label="Обогащённый бор", linewidth=3, color="#d62728")

    for name, thickness_nm in samples:
        y_marker = probability(np.array([thickness_nm]), enriched_enrichment)[0]
        ax.axvline(thickness_nm, color="#666666", linestyle="--", linewidth=1.6, alpha=0.75)
        ax.scatter([thickness_nm], [y_marker], color="#d62728", s=60, zorder=5)
        ax.text(thickness_nm + 20, min(0.96, y_marker + 0.035), name, fontsize=13, rotation=90, va="bottom", ha="left")

    ax.set_xlim(0, 3000)
    ax.set_ylim(0, 1)
    ax.set_xlabel("Толщина покрытия B$_4$C, нм", fontsize=16)
    ax.set_ylabel("Расчётная вероятность захвата нейтрона", fontsize=16)
    ax.legend(fontsize=14, loc="lower right", frameon=True)
    ax.grid(True, color="#D9D9D9", linewidth=0.8)
    ax.text(0.015, 0.96, "P(d) = 1 - exp[-Σd / sin(θ)]", transform=ax.transAxes, fontsize=17, fontweight="bold", va="top")

    note = (
        "Расчёт выполнен по экспоненциальному закону ослабления нейтронного пучка. "
        "График показывает вероятность захвата нейтрона в слое B$_4$C,\n"
        "а не полную эффективность регистрации детектора. Прямые нейтронные измерения исследуемых образцов не проводились."
    )
    fig.text(0.08, 0.03, note, fontsize=13, ha="left", va="bottom")
    fig.tight_layout(rect=(0.04, 0.08, 0.98, 0.98))
    fig.savefig(PNG_PATH, dpi=100, facecolor="white")
    plt.close(fig)


def verify() -> None:
    wb = load_workbook(XLSX_PATH, data_only=False)
    assert wb.sheetnames == ["Параметры", "Расчет", "График"]
    assert wb["Расчет"].max_row >= 302
    assert wb["Расчет"]["B2"].data_type == "f"
    assert wb["Расчет"]["C302"].data_type == "f"
    from PIL import Image

    with Image.open(PNG_PATH) as image:
        assert image.size[0] >= 1920 and image.size[1] >= 1080
    print(XLSX_PATH)
    print(PNG_PATH)


if __name__ == "__main__":
    create_workbook()
    create_png()
    verify()
