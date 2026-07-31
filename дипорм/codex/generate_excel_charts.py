from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "Book1_backup_before_BC50_phase_fix_20260525_235624.xlsx"
OUT = ROOT / "_extracted" / "excel_charts_py"


def configure_font():
    candidates = [
        Path(r"C:\Windows\Fonts\times.ttf"),
        Path(r"C:\Windows\Fonts\timesbd.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            font_manager.fontManager.addfont(str(candidate))
    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "font.size": 12,
            "axes.labelsize": 12,
            "xtick.labelsize": 12,
            "ytick.labelsize": 12,
            "legend.fontsize": 12,
            "axes.linewidth": 0.8,
        }
    )


def numeric(value):
    if value is None:
        return None
    return float(value)


def plot_sheet(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in (6, 8, 10, 12, 14, 16)]
        if values[0] is None and values[3] is None:
            continue
        try:
            rows.append([numeric(v) for v in values])
        except (TypeError, ValueError):
            continue

    x_psi = [r[0] for r in rows if r[0] is not None and r[1] is not None]
    psi = [r[1] for r in rows if r[0] is not None and r[1] is not None]
    x_delta = [r[0] for r in rows if r[0] is not None and r[2] is not None]
    delta = [r[2] for r in rows if r[0] is not None and r[2] is not None]
    x_model_psi = [r[3] for r in rows if r[3] is not None and r[4] is not None]
    model_psi = [r[4] for r in rows if r[3] is not None and r[4] is not None]
    x_model_delta = [r[3] for r in rows if r[3] is not None and r[5] is not None]
    model_delta = [r[5] for r in rows if r[3] is not None and r[5] is not None]

    fig, ax = plt.subplots(figsize=(7.2, 4.6), dpi=220)
    ax.plot(x_psi, psi, color="#1f77b4", linewidth=1.4, label="Ψ")
    ax.plot(x_delta, delta, color="#ff7f0e", linewidth=1.4, linestyle="--", label="Δ")
    ax.plot(x_model_psi, model_psi, color="#2ca02c", linewidth=1.4, linestyle=(0, (4, 4)), label="Модель Ψ")
    ax.plot(x_model_delta, model_delta, color="#17becf", linewidth=1.4, linestyle=(0, (4, 4, 1, 4)), label="Модель Δ")
    ax.set_xlabel("Длина волны, нм", labelpad=8)
    ax.set_ylabel("Ψ, Δ, град.", labelpad=8)
    ax.set_xlim(200, 1700)
    ax.grid(False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(loc="upper left", frameon=False, handlelength=2.8, borderaxespad=0.4)
    fig.tight_layout()
    target = OUT / f"{ws.title}.png"
    fig.savefig(target, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return target


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    configure_font()
    wb = load_workbook(XLSX, data_only=True)
    for ws in wb.worksheets[1:5]:
        print(plot_sheet(ws))


if __name__ == "__main__":
    main()
